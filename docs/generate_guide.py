#!/usr/bin/env python3
"""
Generate the Wealth OS User Guide as both a Word (.docx) and a PDF file.

One content source (the BLOCKS list below) is rendered by two backends:
  - python-docx  -> Wealth-OS-User-Guide.docx
  - fpdf2        -> Wealth-OS-User-Guide.pdf

Run:  python docs/generate_guide.py
Outputs are written to the project root (next to README.md).
"""
import os

# ---------------------------------------------------------------------------
# CONTENT  — list of (kind, text) blocks.
#   h1  = document title          h2 = section       h3 = sub-section
#   p   = paragraph               b  = bullet        n  = numbered step
#   q   = FAQ question            a  = FAQ answer
#   note= callout                 hr = divider
# ---------------------------------------------------------------------------
BLOCKS = [
    ("h1", "Wealth OS — User Guide"),
    ("p",  "Your whole financial world — stocks, funds, property, cars, gold and cash, across every market and currency — in one calm dashboard, with a self-grading AI investing desk."),
    ("note", "Educational software, not financial advice. Every AI call, score, rating and tax estimate is informational only. Always verify with primary sources and a qualified professional before acting."),
    ("hr", ""),

    ("h2", "1. What Wealth OS is (in one minute)"),
    ("p",  "Most people's money lives in 5-10 disconnected places: a home broker, an international broker, a crypto exchange, mutual-fund apps, real estate, gold and several bank accounts. Wealth OS pulls all of it into a single, profile-aware net-worth picture with live prices and live currency conversion."),
    ("p",  "It answers four questions in one place:"),
    ("b",  "What do I own? — holdings, funds, bonds, property, cars, gold and cash."),
    ("b",  "What is it worth right now? — live quotes + FX, with per-asset profit/loss."),
    ("b",  "Why did it move and is it healthy? — real indicators, sentiment, warning signs."),
    ("b",  "What should I consider? — an AI CFO chat and a 6-agent AI Desk that issues plans and grades itself."),

    ("h2", "2. Opening the app"),
    ("n",  "Start the backend (data engine) and the frontend (the website) — see the README's Quick Start. With Docker it is one command; otherwise run the backend on port 8000 and the frontend on port 3000."),
    ("n",  "Open your browser at http://localhost:3000."),
    ("n",  "The first time, a Welcome Setup wizard appears (see the next section). After that the dashboard loads."),
    ("note", "You do NOT need any API keys or paid accounts. The app works out of the box using free public data."),

    ("h2", "3. First-run: the Welcome Setup"),
    ("p",  "On a fresh device the app blurs and shows a mandatory 5-step wizard. You cannot skip it — it creates YOUR profile."),
    ("n",  "Welcome — a short intro."),
    ("n",  "Who is this for? — Personal, Family, or Business."),
    ("n",  "About you — your name (required), email (optional), country, display currency and language."),
    ("n",  "Your style — risk appetite, primary goal, and which asset tracks you care about."),
    ("n",  "Review — confirm, then Finish."),
    ("p",  "When you finish, Wealth OS creates a brand-new profile for you and takes you straight to it. The owner/admin profile is never overwritten, and the last remaining admin profile can never be deleted, so the app always has an owner."),

    ("h2", "4. The top bar (your controls)"),
    ("b",  "Search box — find any stock, ETF, fund or coin. Results open a PREVIEW first; they are not added to your portfolio automatically."),
    ("b",  "Profile switcher — switch between profiles, or open the admin 'Family wealth' view (all profiles combined)."),
    ("b",  "Live toggle — pause/resume automatic price refresh. 'Live' = updating."),
    ("b",  "Refresh (the circular arrow) — fetch fresh prices right now."),
    ("b",  "Lock button (the padlock) — instantly lock the open profile (for a coffee break). The app blurs and asks for the password to resume. It stays online in the background."),
    ("b",  "Language selector — switch the whole interface between 11 languages."),
    ("b",  "Settings (the gear) — theme, currency, density, backups and more."),
    ("b",  "Theme toggle — quick dark/light switch."),

    ("h2", "5. Quick-start checklist (do this first)"),
    ("p",  "A fresh profile is EMPTY on purpose — Wealth OS never invents your money. Most screens stay blank until you add your own data. Do these in order:"),
    ("n",  "Add your holdings — go to Portfolio and add each stock/ETF/coin (symbol, quantity, average price). This fills Mission Control, charts and allocations."),
    ("n",  "Add your cash & banks — in Portfolio, use Add Cash & Cashflow to record bank balances and income/expenses."),
    ("n",  "Add property, cars, gold — go to Properties for anything physical."),
    ("n",  "Add Favorites — this is what powers the AI Desk, Screeners and watchlist. Use the search box or the star on any asset page."),
    ("n",  "Add SIPs / mutual funds — in Funds & Bonds or by importing a CSV in the Brokers tab."),
    ("note", "If a screen looks empty, it almost always means 'you haven't added that type of data yet.' See the Troubleshooting FAQ at the end."),

    ("h2", "6. A tour of every screen"),

    ("h3", "Mission Control"),
    ("p",  "Your home dashboard: total net worth, a live clock and world clocks, market open/closed status, a wealth-projection chart, and allocation/geography/sector/performance charts. It shows zero until you add holdings and cash."),

    ("h3", "About Me"),
    ("p",  "A personal profile page: photo, role, company, headline, bio, focus areas and interests, a milestones timeline, a motto, contact links, and a live wealth-at-a-glance snapshot. Click Edit to fill it in; it starts empty but is always present."),

    ("h3", "Live Markets"),
    ("p",  "Market comparison chart, Market Pulse (breadth/mood), Forex majors, most-active and most-volatile movers, top gainers/losers, global top assets and top coins, plus a per-exchange calendar with live open/closed status."),
    ("p",  "Exchange browser: pick an exchange (NASDAQ, NYSE, NSE India, Crypto...) to list its tradable shares & funds, searchable and paginated, with live prices and one-click drill-down. (Full lists are free for US, NSE India, Crypto and Indian mutual funds; other world exchanges show a curated list of major names — see the API note.)"),

    ("h3", "IPOs"),
    ("p",  "Upcoming and recent IPOs grouped by exchange, with links to the official IPO calendars."),

    ("h3", "Portfolio"),
    ("p",  "The heart of the app. Add, edit and sell holdings. The Sell flow routes proceeds to your cashflow and a chosen bank account. There is a combined Add Cash & Cashflow panel and bank-tagged cash accounts (Axis, ICICI, SBI, HDFC, N26, ...)."),
    ("b",  "To add a holding: enter symbol, quantity and average buy price, then Add holding."),
    ("b",  "Buttons (Add holding / Clear) and the import tools sit on the right of each panel."),

    ("h3", "Properties"),
    ("p",  "Houses, cars, gold and more, in any country/currency. Includes inflation + locality price projection for homes, depreciation curves for cars, rental status, and a dispose (sell/gift) flow."),

    ("h3", "Alerts & Plans"),
    ("p",  "Set price alerts on assets and create savings goals/plans. Triggered alerts show up in notifications."),

    ("h3", "Favorites (important!)"),
    ("p",  "Your watchlist — and the fuel for the AI Desk and Screeners. Three ways to add a favorite:"),
    ("b",  "Type in the top search box and pick a result, then add it."),
    ("b",  "Click the star on any Asset Detail page."),
    ("b",  "Use 'Add favorite' on movers, screeners or Live Markets rows."),
    ("note", "If Favorites is empty, the AI Desk will also be empty — the desk only analyzes the assets you favorite."),

    ("h3", "AI Desk"),
    ("p",  "A transparent 6-agent investing engine that runs over your Favorites, grouped by exchange: Watcher, Analyst, Strategist, Advisor, Tracker and Rater. It scores each favorite, predicts direction, issues BUY/HOLD/TRIM plans (entry, target, stop), and grades its own accuracy over time, auto-tuning what works."),
    ("note", "The AI Desk needs Favorites to do anything. Add a few favorites first, give it a moment to fetch data, and the cards appear. When score, trend and risk-reward align, an asset gets a 'green light' and a one-click 'I bought this - track it' option."),

    ("h3", "Screeners"),
    ("p",  "Bullish/Bearish technical screens with a sentiment gauge, MACD meter, RSI bar, smart-money score and insider-trading data."),

    ("h3", "Compare"),
    ("p",  "Put two or more assets side by side to compare price, performance and key stats."),

    ("h3", "News"),
    ("p",  "Real, de-duplicated article headlines with images and source links. The News Center can scope to your holdings/favorites. (Headlines and links only — open the source for the full article.)"),

    ("h3", "Funds & Bonds"),
    ("p",  "Track mutual funds (with NAV history) and bonds. Indian mutual-fund NAVs come free from AMFI/MFAPI."),

    ("h3", "Asset Detail"),
    ("p",  "Open any asset for a full page: interactive price chart (crosshair, volume, range tabs), a near-real-time price ticker, performance/dividends/price/profile/edit tabs, a technical read-out with sentiment gauge, ratings & warning signs (deep-linked to Morningstar, Simply Wall St, GuruFocus), and real news. Use the star here to favorite it."),

    ("h3", "AI CFO"),
    ("p",  "A conversational analyst grounded in YOUR live portfolio. Ask about cashflow, concentration risk, buy zones, goals, property or tax, and it answers from your real numbers."),

    ("h3", "Brokers"),
    ("p",  "Connection status for Trading 212, IBKR, Zerodha/Coin, Coinbase, Groww, Scalable Capital and MF Central, plus a place to record beneficiary/recovery info (never passwords)."),
    ("note", "Most retail brokers offer no free official sync for outside apps, so use CSV import: drop a SIP or holdings CSV into the Brokers tab and Wealth OS auto-routes it to the right place. A sample SIP CSV is in the 'samples' folder."),

    ("h3", "Profiles"),
    ("p",  "Create one profile per family member, each with its own holdings, property, cash and goals. Click a profile card to edit its name, relation, country, photo and password. You can also choose a different display currency just for this tab."),

    ("h3", "Settings"),
    ("p",  "Theme (dark/light/auto), accent color, density, privacy blur, 12/24h clock, default display currency, auto-refresh interval, startup profile and language. Also includes a one-click privacy-safe JSON backup export (no passwords or API keys)."),

    ("h2", "7. Profiles, passwords & locking"),
    ("b",  "Set a password: Profiles tab -> click a profile card -> type a password in the password field -> Update. A padlock icon appears on protected profiles."),
    ("b",  "Switching: every time you switch INTO a password-protected profile, you must enter its password. Only the profile currently in view stays unlocked."),
    ("b",  "Refresh: protected profiles re-lock on every page refresh - the app stays blurred behind the unlock prompt until you enter the password."),
    ("b",  "Lock now: click the padlock button in the top bar to lock immediately (e.g. stepping away). Enter the password to resume."),
    ("b",  "Family wealth (admin) view: shows everyone combined and requires the admin password."),
    ("note", "The lock secures the on-screen experience for personal/family use. It is an app-level lock (it blurs and gates the UI), not a hardened server login - that is the natural next step if you host this publicly."),

    ("h2", "8. Currencies & languages"),
    ("b",  "Display currency: set it in Settings (or just for the Profiles tab in that tab). All values convert live using real FX rates."),
    ("b",  "Each profile shows in its own home currency by default; you can override it with the currency chips."),
    ("b",  "Language: switch the entire interface between 11 languages from the top bar or Settings, including full right-to-left layout for Arabic."),

    ("h2", "9. Backups (don't skip this)"),
    ("b",  "Settings -> Export backup downloads a private JSON copy of your data (no passwords or API keys)."),
    ("b",  "All your data also lives in one file: backend/data/state.json. Copy it to back up or move to another machine."),
    ("note", "Because there is no cloud database, that one file IS your data. Back it up regularly."),

    ("hr", ""),
    ("h2", "10. Troubleshooting — \"Why can't I see anything?\""),
    ("p",  "Almost every 'empty' screen means the data for it hasn't been added yet. Here are the common ones:"),

    ("q", "The AI Desk is empty / it shows nothing."),
    ("a", "The AI Desk only analyzes your Favorites. Add a few favorites first (search box, or the star on any asset, or 'Add favorite' on a movers/screener row), wait a few seconds for prices to load, then return to AI Desk. If it is still empty, check the Live toggle is on and click Refresh."),

    ("q", "My Favorites list is empty."),
    ("a", "Nothing is favorited yet. Type a company, ticker or coin in the top search box, open the result, and click the star; or use 'Add favorite' buttons on Live Markets / Screeners."),

    ("q", "Mission Control / net worth shows 0."),
    ("a", "You haven't added holdings, cash or property yet. Start with Portfolio -> Add holding and Add Cash & Cashflow, then Properties. The dashboard fills in immediately."),

    ("q", "Prices aren't changing / not truly live."),
    ("a", "Free data is near-real-time, not tick-by-tick: it refreshes on an interval with a small delay. Make sure the Live toggle is ON and press Refresh. True streaming requires a paid feed. US, EU and India stocks/ETFs, all crypto, Indian mutual funds, FX and gold/silver update reliably for free."),

    ("q", "Live price for my specific exchange isn't updating."),
    ("a", "Some exchanges block server-side requests or have no free feed. Keyless live prices work for US (NYSE/Nasdaq), EU (XETRA/LSE/Euronext), India (NSE/BSE) stocks & ETFs, all crypto, Indian mutual funds, FX and metals. Other markets fall back to cached/curated data until a paid key is added."),

    ("q", "The Exchange browser won't list every share for my market."),
    ("a", "Full directories are free only for NASDAQ, NYSE, AMEX, NSE India, Crypto and Indian mutual funds. Frankfurt, LSE, Euronext, Tokyo, Hong Kong and ASX have no free full feed, so a curated list of major companies & ETFs is shown. Adding a paid market-data key unlocks the full lists automatically."),

    ("q", "My broker won't connect / sync."),
    ("a", "Most retail brokers (Trading 212, Groww, MF Central, Scalable, Zerodha) don't offer free official sync for third-party apps. Use CSV/statement import instead: drop your SIP or holdings CSV into the Brokers tab and it auto-routes. See the sample CSV in the 'samples' folder."),

    ("q", "News is empty."),
    ("a", "News is scoped to your assets. Add holdings or favorites and refresh. There is also a general news feed; if it is blank, check your internet connection and press Refresh."),

    ("q", "It keeps asking for my password when I switch profiles."),
    ("a", "That's intended: protected profiles ask for the password every time you switch into them, and re-lock on refresh. Only the profile you're currently viewing stays unlocked. To stop prompts for a profile, remove its password in the Profiles tab."),

    ("q", "I forgot a profile password."),
    ("a", "Passwords are stored in backend/data/state.json on your own machine. If you are the owner, you can edit that file to clear the password field, then restart the backend."),

    ("q", "I changed app.js/css but nothing updated."),
    ("a", "The browser caches the files. Bump the ?v= version on the script/link tags in index.html (developers), or hard-refresh the page (Ctrl/Cmd+Shift+R)."),

    ("q", "Everything looks blank / 'demo' mode."),
    ("a", "The frontend can't reach the backend. Make sure the backend is running on port 8000 and the frontend on port 3000, and that no firewall blocks them."),

    ("h2", "11. Which markets are live for free (summary)"),
    ("b",  "Reliable & keyless: US stocks/ETFs (NYSE, Nasdaq), EU stocks/ETFs (XETRA, LSE, Euronext), India stocks/ETFs (NSE, BSE), ALL crypto, Indian mutual funds (AMFI), FX/forex (ECB), and gold/silver."),
    ("b",  "Curated (not full) without a key: complete share directories for Frankfurt, LSE, Euronext, Tokyo, Hong Kong, ASX."),
    ("b",  "Needs a paid key: tick-by-tick streaming, international full directories, richer company fundamentals, and live broker auto-sync."),

    ("hr", ""),
    ("h2", "Disclaimer"),
    ("p",  "Wealth OS is an educational and personal-finance organization tool. Nothing in it — including AI Desk calls, AI CFO answers, sentiment scores, technical read-outs, warning signs and tax estimates — is financial, investment, legal or tax advice. Markets are risky; verify every figure with primary sources and a qualified professional before acting."),
    ("p",  "Wealth OS — see your whole financial world, clearly."),
]

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
DOCX_PATH = os.path.join(ROOT, "Wealth-OS-User-Guide.docx")
PDF_PATH  = os.path.join(ROOT, "Wealth-OS-User-Guide.pdf")

NAVY = (0x0f, 0x2d, 0x4f)
BLUE = (0x2f, 0x6f, 0xd6)
GREY = (0x44, 0x4a, 0x54)
LIGHT = (0xef, 0xf3, 0xfb)


# --------------------------- DOCX ------------------------------------------
def build_docx():
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    base = doc.styles["Normal"]
    base.font.name = "Calibri"
    base.font.size = Pt(11)

    for kind, text in BLOCKS:
        if kind == "h1":
            p = doc.add_paragraph()
            run = p.add_run(text)
            run.bold = True
            run.font.size = Pt(26)
            run.font.color.rgb = RGBColor(*NAVY)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif kind == "h2":
            p = doc.add_paragraph()
            run = p.add_run(text)
            run.bold = True
            run.font.size = Pt(16)
            run.font.color.rgb = RGBColor(*BLUE)
            p.space_before = Pt(10)
        elif kind == "h3":
            p = doc.add_paragraph()
            run = p.add_run(text)
            run.bold = True
            run.font.size = Pt(13)
            run.font.color.rgb = RGBColor(*NAVY)
        elif kind == "p":
            doc.add_paragraph(text)
        elif kind == "b":
            doc.add_paragraph(text, style="List Bullet")
        elif kind == "n":
            doc.add_paragraph(text, style="List Number")
        elif kind == "q":
            p = doc.add_paragraph()
            run = p.add_run("Q:  " + text)
            run.bold = True
            run.font.color.rgb = RGBColor(*NAVY)
        elif kind == "a":
            p = doc.add_paragraph()
            run = p.add_run("A:  " + text)
            run.font.color.rgb = RGBColor(*GREY)
        elif kind == "note":
            p = doc.add_paragraph()
            run = p.add_run("Note:  " + text)
            run.italic = True
            run.font.color.rgb = RGBColor(*BLUE)
        elif kind == "hr":
            doc.add_paragraph("_" * 60)

    doc.save(DOCX_PATH)
    print("wrote", DOCX_PATH)


# --------------------------- PDF -------------------------------------------
def _ascii(s):
    repl = {
        "—": "-", "–": "-", "‘": "'", "’": "'",
        "“": '"', "”": '"', "…": "...", "→": "->",
        "₹": "Rs", "€": "EUR", "•": "-", " ": " ",
    }
    for k, v in repl.items():
        s = s.replace(k, v)
    return s.encode("latin-1", "ignore").decode("latin-1")


def build_pdf():
    from fpdf import FPDF

    pdf = FPDF(format="A4")
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.set_margins(18, 18, 18)
    pdf.add_page()
    W = pdf.w - pdf.l_margin - pdf.r_margin

    def text_block(txt, size=11, style="", color=(0, 0, 0), gap=2, indent=0, lead=""):
        pdf.set_font("Helvetica", style, size)
        pdf.set_text_color(*color)
        if indent:
            pdf.set_x(pdf.l_margin + indent)
        if lead:
            pdf.set_font("Helvetica", "B", size)
            lead_w = pdf.get_string_width(lead) + 1
            pdf.cell(lead_w, 6, lead)
            pdf.set_font("Helvetica", style, size)
            pdf.multi_cell(W - indent - lead_w, 6, _ascii(txt))
        else:
            pdf.multi_cell(W - indent, 6, _ascii(txt))
        if gap:
            pdf.ln(gap)

    for kind, text in BLOCKS:
        if kind == "h1":
            pdf.set_fill_color(*NAVY)
            pdf.set_text_color(255, 255, 255)
            pdf.set_font("Helvetica", "B", 22)
            pdf.multi_cell(W, 12, _ascii(text), fill=True)
            pdf.ln(3)
        elif kind == "h2":
            pdf.ln(2)
            text_block(text, size=15, style="B", color=BLUE, gap=2)
        elif kind == "h3":
            text_block(text, size=12, style="B", color=NAVY, gap=1)
        elif kind == "p":
            text_block(text, size=11, color=(20, 20, 20))
        elif kind == "b":
            text_block(text, size=11, color=(20, 20, 20), indent=6, lead="-  ", gap=1)
        elif kind == "n":
            text_block(text, size=11, color=(20, 20, 20), indent=6, lead=">  ", gap=1)
        elif kind == "q":
            text_block(text, size=11, style="B", color=NAVY, gap=0, lead="Q:  ")
        elif kind == "a":
            text_block(text, size=11, color=GREY, gap=3, lead="A:  ")
        elif kind == "note":
            y = pdf.get_y()
            pdf.set_fill_color(*LIGHT)
            pdf.set_draw_color(*BLUE)
            pdf.set_text_color(*GREY)
            pdf.set_font("Helvetica", "I", 10)
            pdf.multi_cell(W, 6, _ascii("Note:  " + text), border=0, fill=True)
            pdf.ln(3)
        elif kind == "hr":
            pdf.set_draw_color(180, 190, 205)
            pdf.line(pdf.l_margin, pdf.get_y(), pdf.w - pdf.r_margin, pdf.get_y())
            pdf.ln(4)

    pdf.output(PDF_PATH)
    print("wrote", PDF_PATH)


# --------------------- README.md -> PDF (basic markdown) -------------------
import re

README_PDF = os.path.join(ROOT, "Wealth-OS-README.pdf")
README_MD = os.path.join(ROOT, "README.md")


def _md_inline(s):
    s = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", s)   # links -> text (url)
    s = s.replace("**", "").replace("`", "")
    return s


def build_readme_pdf():
    if not os.path.exists(README_MD):
        return
    from fpdf import FPDF
    pdf = FPDF(format="A4")
    pdf.set_auto_page_break(auto=True, margin=16)
    pdf.set_margins(16, 16, 16)
    pdf.add_page()
    W = pdf.w - pdf.l_margin - pdf.r_margin

    lines = open(README_MD, encoding="utf-8").read().splitlines()
    i = 0
    in_code = False
    in_comment = False
    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        # skip HTML comment blocks (e.g. the screenshot template)
        if in_comment:
            if "-->" in line:
                in_comment = False
            i += 1
            continue
        if line.strip().startswith("<!--"):
            if "-->" not in line:
                in_comment = True
            i += 1
            continue
        if line.strip().startswith("```"):
            in_code = not in_code
            i += 1
            continue
        if in_code:
            pdf.set_font("Courier", "", 8.5)
            pdf.set_text_color(60, 60, 60)
            pdf.multi_cell(W, 4.5, _ascii(line) or " ")
            i += 1
            continue
        # tables
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                rows.append(lines[i])
                i += 1
            for r_idx, row in enumerate(rows):
                cells = [c.strip() for c in row.strip().strip("|").split("|")]
                if cells and set("".join(cells)) <= set("-: "):
                    continue  # separator row
                head = (r_idx == 0)
                first = _md_inline(cells[0]) if cells else ""
                rest = " - ".join(_md_inline(c) for c in cells[1:] if c)
                pdf.set_font("Helvetica", "B" if head else "B", 9.5)
                pdf.set_text_color(*NAVY)
                pdf.multi_cell(W, 5, _ascii(first))
                if rest:
                    pdf.set_font("Helvetica", "", 9.5)
                    pdf.set_text_color(40, 40, 40)
                    pdf.set_x(pdf.l_margin + 5)
                    pdf.multi_cell(W - 5, 5, _ascii(rest))
                pdf.ln(0.6)
            pdf.ln(1)
            continue
        stripped = line.lstrip()
        if stripped.startswith("# "):
            pdf.set_fill_color(*NAVY); pdf.set_text_color(255, 255, 255)
            pdf.set_font("Helvetica", "B", 18)
            pdf.multi_cell(W, 10, _ascii(_md_inline(stripped[2:])), fill=True); pdf.ln(2)
        elif stripped.startswith("## "):
            pdf.ln(1); pdf.set_text_color(*BLUE); pdf.set_font("Helvetica", "B", 14)
            pdf.multi_cell(W, 7, _ascii(_md_inline(stripped[3:]))); pdf.ln(1)
        elif stripped.startswith("### "):
            pdf.set_text_color(*NAVY); pdf.set_font("Helvetica", "B", 11.5)
            pdf.multi_cell(W, 6, _ascii(_md_inline(stripped[4:]))); pdf.ln(0.5)
        elif stripped.startswith("> "):
            pdf.set_fill_color(*LIGHT); pdf.set_text_color(*GREY); pdf.set_font("Helvetica", "I", 10)
            pdf.multi_cell(W, 5.5, _ascii(_md_inline(stripped[2:])), fill=True); pdf.ln(1)
        elif stripped.startswith("- ") or stripped.startswith("* "):
            pdf.set_text_color(30, 30, 30); pdf.set_font("Helvetica", "", 10)
            pdf.set_x(pdf.l_margin + 4)
            pdf.multi_cell(W - 4, 5.5, _ascii("- " + _md_inline(stripped[2:]))); pdf.ln(0.4)
        elif re.match(r"^\d+\.\s", stripped):
            pdf.set_text_color(30, 30, 30); pdf.set_font("Helvetica", "", 10)
            pdf.set_x(pdf.l_margin + 4)
            pdf.multi_cell(W - 4, 5.5, _ascii(_md_inline(stripped))); pdf.ln(0.4)
        elif stripped.startswith("---") or stripped.startswith("==="):
            pdf.set_draw_color(180, 190, 205)
            pdf.line(pdf.l_margin, pdf.get_y(), pdf.w - pdf.r_margin, pdf.get_y()); pdf.ln(3)
        elif stripped == "":
            pdf.ln(2)
        else:
            pdf.set_text_color(25, 25, 25); pdf.set_font("Helvetica", "", 10.5)
            pdf.multi_cell(W, 5.5, _ascii(_md_inline(stripped))); pdf.ln(1)
        i += 1

    pdf.output(README_PDF)
    print("wrote", README_PDF)


if __name__ == "__main__":
    build_docx()
    build_pdf()
    build_readme_pdf()
